import io
import os
import unittest
from datetime import date

os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["DISABLE_BACKGROUND_SCHEDULER"] = "1"

try:
    from openpyxl import load_workbook
    from app import _group_client_records, app
    from models import Cliente, ContabilidadRegistro, Usuario, db
    _IMPORT_ERROR = ""
except ModuleNotFoundError as exc:  # El CI ligero puede omitir dependencias web/PDF.
    app = None
    _IMPORT_ERROR = str(exc)


@unittest.skipIf(app is None, f"Dependencias de integración no disponibles: {_IMPORT_ERROR}")
class ClientesReportesTest(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        with app.app_context():
            db.drop_all()
            db.create_all()
            user = Usuario(
                nombre="clientes_reportes_test",
                nombre_visible="Clientes Reportes Test",
                correo="clientes.reportes@example.com",
                rol="ADMIN",
            )
            user.set_password("Clientes123")
            db.session.add(user)
            db.session.flush()

            db.session.add_all([
                Cliente(
                    nombre_cliente="Constructora Águila",
                    empresa="Águila Norte",
                    responsable="Clientes Reportes Test",
                    razon_social="CONSTRUCTORA AGUILA SA DE CV",
                    rfc="CAA010101AAA",
                    regimen_fiscal="601",
                    codigo_postal_fiscal="64000",
                    uso_cfdi="G03",
                    correo="norte@example.com",
                ),
                Cliente(
                    nombre_cliente="  CONSTRUCTORA AGUILA  ",
                    empresa="Águila Sur",
                    responsable="Clientes Reportes Test",
                    razon_social="AGUILA PROYECTOS SA DE CV",
                    rfc="APS010101AAA",
                    regimen_fiscal="601",
                    codigo_postal_fiscal="66000",
                    uso_cfdi="G03",
                    telefono="8112345678",
                ),
                ContabilidadRegistro(
                    folio="CLI-TEST-001",
                    tipo="CLIENTE",
                    nombre="Constructora Águila",
                    proyecto="Obra Norte",
                    monto_total=1000,
                    moneda="MXN",
                    fecha_inicio=date(2026, 10, 1),
                    fecha_vencimiento=date(2026, 10, 31),
                    tiempo_credito_dias=30,
                    estatus="PENDIENTE",
                    creado_por_id=user.id,
                ),
                ContabilidadRegistro(
                    folio="CLI-TEST-002",
                    tipo="CLIENTE",
                    nombre="CONSTRUCTORA ÁGUILA",
                    proyecto="Obra Sur",
                    monto_total=2000,
                    moneda="MXN",
                    fecha_inicio=date(2026, 10, 2),
                    fecha_vencimiento=date(2026, 11, 1),
                    tiempo_credito_dias=30,
                    estatus="PENDIENTE",
                    creado_por_id=user.id,
                ),
            ])
            db.session.commit()
            self.user_id = user.id

        self.client = app.test_client()
        with self.client.session_transaction() as session:
            session["_user_id"] = str(self.user_id)
            session["_fresh"] = True

    def test_directorio_condensa_por_nombre_y_exporta_excel_pdf(self):
        with app.app_context():
            groups = _group_client_records(Cliente.query.order_by(Cliente.id).all())
            self.assertEqual(len(groups), 1)
            self.assertEqual(groups[0]["record_count"], 2)
            self.assertEqual(len(groups[0]["fiscal"]), 2)

        response = self.client.get("/clientes")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"2 fichas", response.data)
        self.assertIn("Clientes por nombre".encode("utf-8"), response.data)

        excel = self.client.get("/clientes/export.xlsx")
        self.assertEqual(excel.status_code, 200)
        self.assertTrue(excel.data.startswith(b"PK"))
        self.assertIn("spreadsheetml", excel.content_type)

        pdf = self.client.get("/clientes/export.pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b"%PDF"))
        self.assertEqual(pdf.content_type, "application/pdf")

    def test_contabilidad_condensa_clientes_y_exporta_pdf_en_apartados(self):
        response = self.client.get("/contabilidad/clientes")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Saldo por cobrar".encode("utf-8"), response.data)
        self.assertIn("Ver movimientos".encode("utf-8"), response.data)

        excel = self.client.get("/contabilidad/clientes/exportar.xlsx")
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(io.BytesIO(excel.data), read_only=True)
        self.assertEqual(workbook.sheetnames[0], "Resumen por nombre")
        self.assertEqual(workbook["Resumen por nombre"].max_row, 2)

        paths = [
            "/contabilidad/exportar.pdf",
            "/contabilidad/bancos/exportar.pdf",
            "/contabilidad/clientes/exportar.pdf",
            "/contabilidad/proveedores/exportar.pdf",
            "/contabilidad/trabajadores/exportar.pdf",
            "/contabilidad/maquinaria/exportar.pdf",
            "/contabilidad/transporte/exportar.pdf",
        ]
        for path in paths:
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.data.startswith(b"%PDF"))
                self.assertEqual(response.content_type, "application/pdf")


if __name__ == "__main__":
    unittest.main()
