from __future__ import annotations

import io
import json
import os
import secrets
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from flask import (
    Blueprint,
    Response,
    abort,
    current_app,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from werkzeug.utils import secure_filename

from contabilidad_access import can_access_contabilidad, can_manage_contabilidad
from models import ContabilidadAbono, ContabilidadDocumento, ContabilidadRegistro, db


contabilidad_bp = Blueprint("contabilidad", __name__, url_prefix="/contabilidad")

MAX_PDF_BYTES = 15 * 1024 * 1024
MONEY_QUANTUM = Decimal("0.01")

CATEGORIES = {
    "clientes": {
        "tipo": "CLIENTE",
        "singular": "cliente",
        "label": "Clientes",
        "icon": "👥",
        "description": "Créditos y obras por cobrar, abonos, registros y facturas.",
        "financial": True,
        "flow": "cobrar",
    },
    "proveedores": {
        "tipo": "PROVEEDOR",
        "singular": "proveedor",
        "label": "Proveedores",
        "icon": "🏭",
        "description": "Créditos por pagar, abonos, expedientes y facturas.",
        "financial": True,
        "flow": "pagar",
    },
    "trabajadores": {
        "tipo": "TRABAJADOR",
        "singular": "trabajador o empleado",
        "label": "Trabajadores y empleados",
        "icon": "👷",
        "description": "Altas de personal por proyecto y documentación requerida.",
        "financial": False,
        "flow": "administrar",
    },
    "maquinaria": {
        "tipo": "MAQUINARIA",
        "singular": "equipo o maquinaria",
        "label": "Maquinaria y equipo",
        "icon": "🏗️",
        "description": "Padrón de maquinaria, equipo y sus expedientes.",
        "financial": False,
        "flow": "administrar",
    },
    "transporte": {
        "tipo": "TRANSPORTE",
        "singular": "equipo de transporte",
        "label": "Equipo de transporte",
        "icon": "🚚",
        "description": "Control de unidades, placas y documentación.",
        "financial": False,
        "flow": "administrar",
    },
}

CATEGORY_BY_TYPE = {item["tipo"]: dict(slug=slug, **item) for slug, item in CATEGORIES.items()}

DOCUMENT_TYPES = {
    "REGISTRO": "Registro / alta",
    "FACTURA": "Factura",
    "DOCUMENTACION": "Documentación general",
    "CONTRATO": "Contrato",
    "IDENTIFICACION": "Identificación",
    "FISCAL": "Documento fiscal",
    "SEGURO": "Seguro / póliza",
    "TARJETA_CIRCULACION": "Tarjeta de circulación",
}


@contabilidad_bp.before_request
def _restrict_accounting_module():
    if getattr(current_user, "is_authenticated", False) and not can_access_contabilidad(current_user):
        abort(403)


def _category_or_404(slug: str) -> dict:
    category = CATEGORIES.get((slug or "").strip().lower())
    if not category:
        abort(404)
    return dict(slug=slug, **category)


def _record_or_404(slug: str, record_id: int) -> tuple[dict, ContabilidadRegistro]:
    category = _category_or_404(slug)
    record = db.session.get(ContabilidadRegistro, record_id)
    if record is None or record.tipo != category["tipo"]:
        abort(404)
    return category, record


def _current_user_name() -> str:
    return (
        getattr(current_user, "nombre_visible", None)
        or getattr(current_user, "nombre", None)
        or "Usuario"
    ).strip()


def _money(raw: str | None, *, required: bool = False) -> float:
    normalized = (raw or "").strip().replace(",", "")
    if not normalized:
        if required:
            raise ValueError("Captura un monto válido.")
        return 0.0
    try:
        value = Decimal(normalized).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ValueError("Captura un monto válido.") from exc
    if value < 0:
        raise ValueError("El monto no puede ser negativo.")
    return float(value)


def _date(raw: str | None, *, required: bool = True):
    value = (raw or "").strip()
    if not value and not required:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError("Captura una fecha válida.") from exc


def _optional_year(raw: str | None) -> int | None:
    value = (raw or "").strip()
    if not value:
        return None
    try:
        year = int(value)
    except ValueError as exc:
        raise ValueError("El año debe ser numérico.") from exc
    if year < 1900 or year > 2200:
        raise ValueError("El año debe estar entre 1900 y 2200.")
    return year


def _credit_days(raw: str | None) -> int | None:
    value = (raw or "").strip()
    if not value:
        return None
    try:
        days = int(value)
    except ValueError as exc:
        raise ValueError("El tiempo de crédito debe indicarse en días.") from exc
    if days < 0 or days > 3650:
        raise ValueError("El tiempo de crédito debe estar entre 0 y 3650 días.")
    return days


def _can_manage_records() -> bool:
    return can_manage_contabilidad(current_user)


def _require_manage() -> None:
    if not _can_manage_records():
        abort(403)


def _record_payload(record: ContabilidadRegistro) -> dict:
    detail_parts = []
    if record.razon_social and record.razon_social.casefold() != record.nombre.casefold():
        detail_parts.append(record.razon_social)
    if record.rfc:
        detail_parts.append(f"RFC {record.rfc}")
    if record.identificador:
        detail_parts.append(record.identificador)
    if record.marca or record.modelo:
        detail_parts.append(" ".join(item for item in (record.marca, record.modelo) if item))
    label = " · ".join([record.nombre, *detail_parts])
    return {
        "id": record.id,
        "folio": record.folio,
        "label": label,
        "nombre": (record.nombre or "").strip(),
        "razon_social": (record.razon_social or "").strip(),
        "rfc": (record.rfc or "").strip(),
        "regimen_fiscal": (record.regimen_fiscal or "").strip(),
        "codigo_postal_fiscal": (record.codigo_postal_fiscal or "").strip(),
        "uso_cfdi": (record.uso_cfdi or "").strip(),
        "contacto": (record.contacto or "").strip(),
        "correo": (record.correo or "").strip(),
        "telefono": (record.telefono or "").strip(),
        "direccion": (record.direccion or "").strip(),
        "identificador": (record.identificador or "").strip(),
        "marca": (record.marca or "").strip(),
        "modelo": (record.modelo or "").strip(),
        "anio": record.anio or "",
        "descripcion": (record.descripcion or "").strip(),
    }


def _record_identity(record: ContabilidadRegistro) -> str:
    if record.tipo in {"CLIENTE", "PROVEEDOR"}:
        value = record.rfc or record.razon_social or record.nombre
    else:
        value = record.identificador or record.nombre
    return (value or str(record.id)).strip().casefold()


def _load_registered_records(category: dict, *, exclude_id: int | None = None) -> list[dict]:
    """Devuelve la versión más reciente de cada entidad registrada en la categoría."""
    records = (
        ContabilidadRegistro.query
        .filter(ContabilidadRegistro.tipo == category["tipo"])
        .order_by(ContabilidadRegistro.actualizado_en.desc(), ContabilidadRegistro.id.desc())
        .all()
    )
    seen: set[str] = set()
    result: list[dict] = []
    for record in records:
        if exclude_id and record.id == exclude_id:
            continue
        identity = _record_identity(record)
        if identity in seen:
            continue
        seen.add(identity)
        result.append(_record_payload(record))
    return sorted(result, key=lambda item: item["label"].casefold())


def _load_altas(category: dict) -> list[dict]:
    """Lee el catálogo histórico de Altas y lo integra como fuente de captura."""
    path = Path(current_app.root_path) / "provider_numbers.json"
    try:
        raw_rows = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return []
    if not isinstance(raw_rows, list):
        return []

    expected_relation = category["tipo"]
    rows: list[dict] = []
    for position, raw in enumerate(raw_rows, start=1):
        if not isinstance(raw, dict):
            continue
        relation = str(raw.get("relacion") or "PROVEEDOR").strip().upper()
        if relation != expected_relation:
            continue
        rows.append(
            {
                "id": position,
                "numero": str(raw.get("numero") or "").strip(),
                "empresa": str(raw.get("empresa") or "").strip(),
                "razon_social": str(raw.get("razon_social_poliutech") or "").strip(),
                "contacto": str(raw.get("contacto") or "").strip(),
                "telefono": str(raw.get("telefono") or "").strip(),
                "correo": str(raw.get("correo") or "").strip(),
                "credito": bool(raw.get("credito", False)),
                "monto_credito": str(raw.get("monto_credito") or "").strip(),
            }
        )
    return rows


def _next_folio(category: dict) -> str:
    prefixes = {
        "CLIENTE": "CC",
        "PROVEEDOR": "CP",
        "TRABAJADOR": "CT",
        "MAQUINARIA": "CM",
        "TRANSPORTE": "CV",
    }
    prefix = prefixes[category["tipo"]]
    return f"{prefix}-{datetime.now().strftime('%Y%m%d%H%M%S')}-{secrets.token_hex(2).upper()}"


def _upload_root() -> Path:
    configured = (os.getenv("UPLOAD_STORAGE_ROOT") or "").strip()
    if configured:
        root = Path(configured).expanduser().resolve()
    elif Path("/data").is_dir():
        root = Path("/data/uploads")
    else:
        root = (Path(current_app.root_path) / "static" / "uploads").resolve()
    return root / "contabilidad"


def _document_path(document: ContabilidadDocumento) -> Path:
    root = _upload_root().resolve()
    candidate = (root / (document.ruta or "")).resolve()
    try:
        candidate.relative_to(root)
    except ValueError:
        abort(404)
    return candidate


def _save_pdf(upload, record: ContabilidadRegistro, document_type: str, description: str = "") -> ContabilidadDocumento | None:
    if not upload or not (upload.filename or "").strip():
        return None
    original = secure_filename(upload.filename) or "documento.pdf"
    if Path(original).suffix.lower() != ".pdf":
        raise ValueError(f"{upload.filename}: solamente se permiten archivos PDF.")
    content = upload.read(MAX_PDF_BYTES + 1)
    if not content:
        raise ValueError(f"{upload.filename}: el archivo está vacío.")
    if len(content) > MAX_PDF_BYTES:
        raise ValueError(f"{upload.filename}: el PDF excede el límite de 15 MB.")
    if not content.lstrip().startswith(b"%PDF-"):
        raise ValueError(f"{upload.filename}: el archivo no parece ser un PDF válido.")

    record_directory = _upload_root() / str(record.id)
    record_directory.mkdir(parents=True, exist_ok=True)
    stored_name = f"{secrets.token_hex(16)}.pdf"
    disk_path = record_directory / stored_name
    disk_path.write_bytes(content)
    return ContabilidadDocumento(
        registro_id=record.id,
        tipo=document_type,
        nombre_original=original,
        nombre_archivo=stored_name,
        ruta=f"{record.id}/{stored_name}",
        mime_type="application/pdf",
        tamano=len(content),
        descripcion=(description or "").strip()[:260] or None,
        usuario_id=getattr(current_user, "id", None),
        usuario_nombre=_current_user_name(),
    )


def _files_from_initial_form(category: dict) -> list[tuple[object, str, str]]:
    uploads: list[tuple[object, str, str]] = []
    if category["tipo"] in {"CLIENTE", "PROVEEDOR"}:
        registration = request.files.get("registro_pdf")
        if registration and registration.filename:
            uploads.append((registration, "REGISTRO", "Registro inicial"))
        uploads.extend((item, "FACTURA", "Factura inicial") for item in request.files.getlist("facturas") if item.filename)
    else:
        uploads.extend((item, "DOCUMENTACION", "Documento de alta") for item in request.files.getlist("documentos") if item.filename)
    return uploads


def _apply_form(record: ContabilidadRegistro, category: dict) -> None:
    source_record = None
    source_id = request.form.get("registro_origen_id", type=int)
    if source_id:
        source_record = db.session.get(ContabilidadRegistro, source_id)
        if source_record is None or source_record.tipo != category["tipo"] or source_record.id == record.id:
            raise ValueError("El registro anterior seleccionado no es válido para esta categoría.")
        if category["tipo"] == "CLIENTE":
            record.cliente_id = source_record.cliente_id

    def _form_value(field: str) -> str:
        submitted = (request.form.get(field) or "").strip()
        if submitted or source_record is None:
            return submitted
        return str(getattr(source_record, field, "") or "").strip()

    name = _form_value("nombre")
    if not name:
        raise ValueError(f"El nombre del {category['singular']} es obligatorio.")
    amount = _money(request.form.get("monto_total"), required=category["financial"])
    if category["financial"] and amount <= 0:
        raise ValueError("El monto del crédito u obra debe ser mayor a cero.")
    if record.id and amount + 0.005 < record.total_abonado:
        raise ValueError("El monto total no puede ser menor que los abonos ya registrados.")

    start_date = _date(request.form.get("fecha_inicio"))
    record.nombre = name[:180]
    record.razon_social = _form_value("razon_social")[:200] or None
    record.rfc = _form_value("rfc").upper()[:20] or None
    record.regimen_fiscal = _form_value("regimen_fiscal")[:10] or None
    record.codigo_postal_fiscal = _form_value("codigo_postal_fiscal")[:10] or None
    record.uso_cfdi = _form_value("uso_cfdi")[:10].upper() or None
    record.contacto = _form_value("contacto")[:160] or None
    record.correo = _form_value("correo")[:160] or None
    record.telefono = _form_value("telefono")[:60] or None
    record.direccion = _form_value("direccion")[:300] or None
    record.proyecto = (request.form.get("proyecto") or "").strip()[:200] or None
    record.identificador = _form_value("identificador")[:100] or None
    record.marca = _form_value("marca")[:100] or None
    record.modelo = _form_value("modelo")[:100] or None
    record.anio = _optional_year(_form_value("anio"))
    record.descripcion = _form_value("descripcion") or None
    record.folio_factura = (request.form.get("folio_factura") or "").strip()[:120] or None
    record.monto_total = amount
    record.moneda = (request.form.get("moneda") or "MXN").strip().upper()[:10] or "MXN"
    record.fecha_inicio = start_date
    record.notas = (request.form.get("notas") or "").strip() or None
    if category["financial"]:
        credit_days = _credit_days(request.form.get("tiempo_credito_dias"))
        due_date = _date(request.form.get("fecha_vencimiento"), required=False)
        if due_date is None and credit_days is None:
            raise ValueError("Captura el tiempo de crédito o la fecha de vencimiento.")
        if due_date is None:
            due_date = start_date + timedelta(days=credit_days or 0)
        if due_date < start_date:
            raise ValueError("La fecha de vencimiento no puede ser anterior a la fecha de inicio.")
        record.fecha_vencimiento = due_date
        record.tiempo_credito_dias = (due_date - start_date).days
        record.estatus = "LIQUIDADO" if record.esta_liquidado and amount > 0 else "PENDIENTE"
    else:
        record.fecha_vencimiento = None
        record.tiempo_credito_dias = None
        record.estatus = "ACTIVO"


def _filtered_query(category: dict | None = None):
    query = ContabilidadRegistro.query
    if category:
        query = query.filter(ContabilidadRegistro.tipo == category["tipo"])
    q = (request.args.get("q") or "").strip()
    status = (request.args.get("estatus") or "").strip().upper()
    if q:
        like = f"%{q}%"
        query = query.filter(
            or_(
                ContabilidadRegistro.folio.ilike(like),
                ContabilidadRegistro.nombre.ilike(like),
                ContabilidadRegistro.razon_social.ilike(like),
                ContabilidadRegistro.rfc.ilike(like),
                ContabilidadRegistro.proyecto.ilike(like),
                ContabilidadRegistro.identificador.ilike(like),
                ContabilidadRegistro.folio_factura.ilike(like),
            )
        )
    if status in {"PENDIENTE", "LIQUIDADO", "ACTIVO", "INACTIVO"}:
        query = query.filter(ContabilidadRegistro.estatus == status)
    return query, q, status


@contabilidad_bp.get("/")
@login_required
def index():
    records = ContabilidadRegistro.query.order_by(ContabilidadRegistro.actualizado_en.desc()).all()
    cards = []
    for slug, data in CATEGORIES.items():
        category_records = [item for item in records if item.tipo == data["tipo"]]
        cards.append(
            {
                **data,
                "slug": slug,
                "count": len(category_records),
                "monto": sum(float(item.monto_total or 0) for item in category_records),
                "saldo": sum(float(item.saldo_pendiente or 0) for item in category_records) if data["financial"] else 0,
            }
        )
    financial = [item for item in records if item.tipo in {"CLIENTE", "PROVEEDOR"}]
    return render_template(
        "contabilidad/index.html",
        cards=cards,
        altas_count=len(_load_altas({"tipo": "CLIENTE"})) + len(_load_altas({"tipo": "PROVEEDOR"})),
        recent=records[:8],
        total_monto=sum(float(item.monto_total or 0) for item in financial),
        total_abonado=sum(float(item.total_abonado or 0) for item in financial),
        total_saldo=sum(float(item.saldo_pendiente or 0) for item in financial),
        category_by_type=CATEGORY_BY_TYPE,
    )


@contabilidad_bp.route("/<slug>", methods=["GET", "POST"])
@login_required
def registros(slug: str):
    category = _category_or_404(slug)
    if request.method == "POST":
        _require_manage()
        record = ContabilidadRegistro(
            folio=_next_folio(category),
            tipo=category["tipo"],
            creado_por_id=getattr(current_user, "id", None),
            creado_por_nombre=_current_user_name(),
        )
        written_paths: list[Path] = []
        try:
            _apply_form(record, category)
            db.session.add(record)
            db.session.flush()
            for upload, document_type, description in _files_from_initial_form(category):
                document = _save_pdf(upload, record, document_type, description)
                if document:
                    db.session.add(document)
                    written_paths.append(_document_path(document))
            db.session.commit()
        except (ValueError, OSError) as exc:
            db.session.rollback()
            for path in written_paths:
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
            flash(str(exc), "warning")
            return redirect(url_for("contabilidad.registros", slug=slug))
        success_label = "Movimiento" if category["financial"] else "Alta"
        flash(f"{success_label} de {category['singular']} registrado correctamente.", "success")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id))

    query, q, status = _filtered_query(category)
    records = query.order_by(ContabilidadRegistro.fecha_inicio.desc(), ContabilidadRegistro.id.desc()).all()
    altas = _load_altas(category) if category["financial"] else []
    selected_alta = None
    try:
        selected_alta_id = int(request.args.get("alta") or 0)
    except (TypeError, ValueError):
        selected_alta_id = 0
    if selected_alta_id:
        selected_alta = next((item for item in altas if item["id"] == selected_alta_id), None)
    return render_template(
        "contabilidad/registros.html",
        category=category,
        records=records,
        altas=altas,
        registered_records=_load_registered_records(category),
        selected_alta=selected_alta,
        can_manage=_can_manage_records(),
        q=q,
        status=status,
        today=datetime.now().date().isoformat(),
    )


@contabilidad_bp.get("/<slug>/<int:record_id>")
@login_required
def detalle(slug: str, record_id: int):
    category, record = _record_or_404(slug, record_id)
    documents_by_type = {
        key: [doc for doc in record.documentos if doc.tipo == key]
        for key in DOCUMENT_TYPES
    }
    return render_template(
        "contabilidad/detalle.html",
        category=category,
        record=record,
        document_types=DOCUMENT_TYPES,
        documents_by_type=documents_by_type,
        registered_records=_load_registered_records(category, exclude_id=record.id),
        can_manage=_can_manage_records(),
        today=datetime.now().date().isoformat(),
    )


@contabilidad_bp.post("/<slug>/<int:record_id>/editar")
@login_required
def editar(slug: str, record_id: int):
    category, record = _record_or_404(slug, record_id)
    _require_manage()
    try:
        _apply_form(record, category)
        db.session.commit()
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#datos")
    flash("Los datos del expediente se actualizaron.", "success")
    return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id))


@contabilidad_bp.post("/<slug>/<int:record_id>/eliminar")
@login_required
def eliminar_registro(slug: str, record_id: int):
    category, record = _record_or_404(slug, record_id)
    _require_manage()
    document_paths = [_document_path(document) for document in (record.documentos or [])]
    record_name = record.nombre
    db.session.delete(record)
    db.session.commit()
    for path in document_paths:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            current_app.logger.warning("No se pudo retirar el PDF contable %s", path)
    flash(f"El registro de {record_name} fue eliminado.", "success")
    return redirect(url_for("contabilidad.registros", slug=category["slug"]))


@contabilidad_bp.post("/<slug>/<int:record_id>/abonos")
@login_required
def agregar_abono(slug: str, record_id: int):
    category, record = _record_or_404(slug, record_id)
    _require_manage()
    if not category["financial"]:
        abort(404)
    try:
        amount = _money(request.form.get("monto"), required=True)
        payment_date = _date(request.form.get("fecha"))
        if amount <= 0:
            raise ValueError("El abono debe ser mayor a cero.")
        if record.esta_liquidado:
            raise ValueError("Este expediente ya está liquidado.")
        if amount - record.saldo_pendiente > 0.005:
            raise ValueError(f"El abono no puede superar el saldo pendiente de ${record.saldo_pendiente:,.2f}.")
        remaining_after_payment = round(max(record.saldo_pendiente - amount, 0.0), 2)
        payment = ContabilidadAbono(
            registro_id=record.id,
            fecha=payment_date,
            monto=amount,
            referencia=(request.form.get("referencia") or "").strip()[:120] or None,
            metodo_pago=(request.form.get("metodo_pago") or "").strip()[:80] or None,
            notas=(request.form.get("notas") or "").strip() or None,
            usuario_id=getattr(current_user, "id", None),
            usuario_nombre=_current_user_name(),
        )
        db.session.add(payment)
        db.session.flush()
        record.estatus = "LIQUIDADO" if remaining_after_payment <= 0.005 else "PENDIENTE"
        db.session.commit()
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#abonos")
    flash("Abono registrado correctamente.", "success")
    return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#abonos")


def _payment_or_404(record: ContabilidadRegistro, payment_id: int) -> ContabilidadAbono:
    payment = db.session.get(ContabilidadAbono, payment_id)
    if payment is None or payment.registro_id != record.id:
        abort(404)
    return payment


@contabilidad_bp.post("/<slug>/<int:record_id>/abonos/<int:payment_id>/editar")
@login_required
def editar_abono(slug: str, record_id: int, payment_id: int):
    category, record = _record_or_404(slug, record_id)
    _require_manage()
    if not category["financial"]:
        abort(404)
    payment = _payment_or_404(record, payment_id)
    try:
        amount = _money(request.form.get("monto"), required=True)
        payment_date = _date(request.form.get("fecha"))
        available = round(record.saldo_pendiente + float(payment.monto or 0), 2)
        if amount <= 0:
            raise ValueError("El abono debe ser mayor a cero.")
        if amount - available > 0.005:
            raise ValueError(f"El abono no puede superar el saldo disponible de ${available:,.2f}.")
        payment.monto = amount
        payment.fecha = payment_date
        payment.referencia = (request.form.get("referencia") or "").strip()[:120] or None
        payment.metodo_pago = (request.form.get("metodo_pago") or "").strip()[:80] or None
        payment.notas = (request.form.get("notas") or "").strip() or None
        total_after_update = sum(float(item.monto or 0) for item in record.abonos)
        record.estatus = "LIQUIDADO" if float(record.monto_total or 0) - total_after_update <= 0.005 else "PENDIENTE"
        db.session.commit()
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + f"#abono-{payment.id}")
    flash("Abono actualizado correctamente.", "success")
    return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#abonos")


@contabilidad_bp.post("/<slug>/<int:record_id>/abonos/<int:payment_id>/eliminar")
@login_required
def eliminar_abono(slug: str, record_id: int, payment_id: int):
    category, record = _record_or_404(slug, record_id)
    _require_manage()
    if not category["financial"]:
        abort(404)
    payment = _payment_or_404(record, payment_id)
    remaining_paid = sum(
        float(item.monto or 0)
        for item in record.abonos
        if item.id != payment.id
    )
    db.session.delete(payment)
    record.estatus = "LIQUIDADO" if float(record.monto_total or 0) - remaining_paid <= 0.005 else "PENDIENTE"
    db.session.commit()
    flash("Abono eliminado correctamente.", "success")
    return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#abonos")


@contabilidad_bp.post("/<slug>/<int:record_id>/documentos")
@login_required
def agregar_documentos(slug: str, record_id: int):
    _, record = _record_or_404(slug, record_id)
    _require_manage()
    document_type = (request.form.get("tipo") or "DOCUMENTACION").strip().upper()
    if document_type not in DOCUMENT_TYPES:
        flash("Selecciona un tipo de documento válido.", "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#documentos")
    uploads = [item for item in request.files.getlist("archivos") if (item.filename or "").strip()]
    if not uploads:
        flash("Selecciona al menos un archivo PDF.", "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#documentos")

    written_paths: list[Path] = []
    try:
        for upload in uploads:
            document = _save_pdf(upload, record, document_type, request.form.get("descripcion") or "")
            if document:
                db.session.add(document)
                written_paths.append(_document_path(document))
        db.session.commit()
    except (ValueError, OSError) as exc:
        db.session.rollback()
        for path in written_paths:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        flash(str(exc), "warning")
        return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#documentos")
    flash(f"Se cargaron {len(written_paths)} documento(s) PDF.", "success")
    return redirect(url_for("contabilidad.detalle", slug=slug, record_id=record.id) + "#documentos")


@contabilidad_bp.get("/documentos/<int:document_id>")
@login_required
def ver_documento(document_id: int):
    document = db.session.get(ContabilidadDocumento, document_id)
    if document is None:
        abort(404)
    path = _document_path(document)
    if not path.is_file():
        abort(404)
    return send_file(
        path,
        mimetype="application/pdf",
        as_attachment=False,
        download_name=document.nombre_original or path.name,
    )


@contabilidad_bp.post("/documentos/<int:document_id>/eliminar")
@login_required
def eliminar_documento(document_id: int):
    _require_manage()
    document = db.session.get(ContabilidadDocumento, document_id)
    if document is None:
        abort(404)
    record = document.registro
    category = CATEGORY_BY_TYPE.get(record.tipo)
    if not category:
        abort(404)
    path = _document_path(document)
    db.session.delete(document)
    db.session.commit()
    try:
        path.unlink(missing_ok=True)
    except OSError:
        current_app.logger.warning("No se pudo retirar el PDF contable %s", path)
    flash("Documento eliminado del expediente.", "success")
    return redirect(url_for("contabilidad.detalle", slug=category["slug"], record_id=record.id) + "#documentos")


@contabilidad_bp.post("/documentos/<int:document_id>/editar")
@login_required
def editar_documento(document_id: int):
    _require_manage()
    document = db.session.get(ContabilidadDocumento, document_id)
    if document is None:
        abort(404)
    record = document.registro
    category = CATEGORY_BY_TYPE.get(record.tipo)
    if not category:
        abort(404)
    document_type = (request.form.get("tipo") or "").strip().upper()
    display_name = Path((request.form.get("nombre_original") or "").strip()).name[:260]
    if document_type not in DOCUMENT_TYPES:
        flash("Selecciona un tipo de documento válido.", "warning")
        return redirect(url_for("contabilidad.detalle", slug=category["slug"], record_id=record.id) + "#documentos")
    if not display_name:
        flash("El nombre visible del documento es obligatorio.", "warning")
        return redirect(url_for("contabilidad.detalle", slug=category["slug"], record_id=record.id) + "#documentos")
    document.tipo = document_type
    document.nombre_original = display_name
    document.descripcion = (request.form.get("descripcion") or "").strip()[:260] or None
    db.session.commit()
    flash("Documento actualizado correctamente.", "success")
    return redirect(url_for("contabilidad.detalle", slug=category["slug"], record_id=record.id) + "#documentos")


def _excel_response(records: list[ContabilidadRegistro], filename_prefix: str) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte contable"
    headers = [
        "Folio",
        "Tipo",
        "Nombre / razón social",
        "RFC",
        "Contacto",
        "Correo",
        "Teléfono",
        "Proyecto",
        "Identificador",
        "Marca",
        "Modelo",
        "Año",
        "Fecha de inicio",
        "Monto total",
        "Total abonado",
        "Saldo pendiente",
        "Moneda",
        "Estatus",
        "Documentos PDF",
        "Notas",
    ]
    ws.append(headers)
    for record in records:
        ws.append(
            [
                record.folio,
                CATEGORY_BY_TYPE.get(record.tipo, {}).get("label", record.tipo),
                record.razon_social or record.nombre,
                record.rfc or "",
                record.contacto or "",
                record.correo or "",
                record.telefono or "",
                record.proyecto or "",
                record.identificador or "",
                record.marca or "",
                record.modelo or "",
                record.anio or "",
                record.fecha_inicio,
                float(record.monto_total or 0),
                float(record.total_abonado or 0),
                float(record.saldo_pendiente or 0),
                record.moneda or "MXN",
                record.estatus,
                len(record.documentos or []),
                record.notas or "",
            ]
        )

    blue_fill = PatternFill("solid", fgColor="0C3C78")
    thin = Side(style="thin", color="D9E2F1")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        row[12].number_format = "dd/mm/yyyy"
        for idx in (13, 14, 15):
            row[idx].number_format = '"$"#,##0.00'
    widths = [24, 24, 34, 18, 24, 28, 18, 28, 20, 16, 16, 10, 16, 18, 18, 18, 12, 16, 16, 42]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 28

    summary = wb.create_sheet("Resumen")
    summary.append(["Concepto", "Valor"])
    summary_rows = [
        ("Registros exportados", len(records)),
        ("Monto total", sum(float(item.monto_total or 0) for item in records)),
        ("Total abonado", sum(float(item.total_abonado or 0) for item in records)),
        ("Saldo pendiente", sum(float(item.saldo_pendiente or 0) for item in records)),
        ("Generado el", datetime.now()),
    ]
    for row in summary_rows:
        summary.append(row)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = blue_fill
    for row_number in (3, 4, 5):
        summary.cell(row=row_number, column=2).number_format = '"$"#,##0.00'
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 24

    output = io.BytesIO()
    wb.save(output)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename_prefix}_{stamp}.xlsx"'},
    )


def _financial_excel_response(records: list[ContabilidadRegistro], filename_prefix: str, sheet_name: str) -> Response:
    """Genera el formato solicitado para los reportes de clientes y proveedores."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    headers = [
        "Nombre / razón social",
        "RFC",
        "Proyecto",
        "FOLIO FACTURA",
        "Fecha de inicio",
        "fecha de vencimiento",
        "Monto total",
        "Total abonado",
        "Saldo pendiente",
        "Estatus",
        "Notas",
    ]
    ws.append(headers)
    for record in records:
        ws.append(
            [
                record.razon_social or record.nombre,
                record.rfc or "",
                record.proyecto or "",
                record.folio_factura or "",
                record.fecha_inicio,
                record.fecha_vencimiento,
                float(record.monto_total or 0),
                float(record.total_abonado or 0),
                float(record.saldo_pendiente or 0),
                record.estatus,
                record.notas or "",
            ]
        )

    header_fill = PatternFill("solid", fgColor="0C3C78")
    separator = Side(style="thin", color="FFFFFF")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=separator, right=separator)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {1, 3, 11})
        for index in (5, 6):
            row[index - 1].number_format = "dd/mm/yyyy"
            row[index - 1].alignment = Alignment(horizontal="center", vertical="top")
        for index in (7, 8, 9):
            row[index - 1].number_format = '"$"#,##0.00'
            row[index - 1].alignment = Alignment(horizontal="right", vertical="top")

    widths = [34, 18, 28, 20, 18, 21, 18, 18, 18, 16, 42]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{max(ws.max_row, 1)}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    output = io.BytesIO()
    wb.save(output)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename_prefix}_{stamp}.xlsx"'},
    )


@contabilidad_bp.get("/exportar.xlsx")
@login_required
def exportar_todo():
    query, _, _ = _filtered_query()
    records = query.order_by(ContabilidadRegistro.tipo.asc(), ContabilidadRegistro.fecha_inicio.desc()).all()
    return _excel_response(records, "contabilidad_general")


@contabilidad_bp.get("/<slug>/exportar.xlsx")
@login_required
def exportar(slug: str):
    category = _category_or_404(slug)
    query, _, _ = _filtered_query(category)
    records = query.order_by(ContabilidadRegistro.fecha_inicio.desc(), ContabilidadRegistro.id.desc()).all()
    if category["financial"]:
        return _financial_excel_response(records, f"reporte_{slug}", category["label"])
    return _excel_response(records, f"contabilidad_{slug}")
